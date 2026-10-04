"""Regression tests for configuration, process budgets and resource boundaries."""

import json
import logging
import subprocess
import sys
from importlib.resources import files
from pathlib import Path
from unittest.mock import patch

import pytest
from hypothesis import given
from hypothesis import strategies as st
from jsonschema import validate

from markdown2docx import MarkdownToDocxConfig, MarkdownToDocxConverter
from markdown2docx.archive import validate_archive
from markdown2docx.config import LimitsConfig, load_config
from markdown2docx.exceptions import (
    ConfigurationError,
    ConversionError,
    ValidationError,
)
from markdown2docx.multilingual import fragments
from markdown2docx.report import ConversionReport
from markdown2docx.resources import validate_resources, validate_restricted_arguments
from markdown2docx.runner import PandocTimeoutError, run_pandoc


@pytest.mark.parametrize(
    "data",
    [
        {"unknown": {}},
        {"pandoc": {"timeout_seconds": -1}},
        {"pandoc": {"timeout_seconds": True}},
        {"pandoc": {"unexpected": 1}},
        {"template": {"margin_cm": "invalid"}},
        {"template": {"margin_cm": float("nan")}},
        {"international": {"offline": "false"}},
        {"limits": {"input_bytes": 0}},
    ],
)
def test_invalid_configuration_fails_early(data):
    with pytest.raises(ConfigurationError) as error:
        MarkdownToDocxConfig.from_dict(data)
    assert error.value.config_key


def test_missing_explicit_config_is_an_error(tmp_path):
    with pytest.raises(ConfigurationError):
        load_config(tmp_path / "missing.toml")


def test_environment_parsing_uses_field_types(monkeypatch):
    monkeypatch.setenv("MD2DOCX_PANDOC__TIMEOUT_SECONDS", "1")
    monkeypatch.setenv("MD2DOCX_TEMPLATE__BODY_FONT", "1")
    config = load_config()
    assert type(config.pandoc.timeout_seconds) is int
    assert config.template.body_font == "1"


def test_config_instances_do_not_share_state():
    with patch.object(MarkdownToDocxConverter, "_validate_pandoc"):
        first, second = MarkdownToDocxConverter(), MarkdownToDocxConverter()
        first.config.template.body_size_pt = 17
        assert second.config.template.body_size_pt == 11
        provided = MarkdownToDocxConfig()
        third = MarkdownToDocxConverter(config=provided)
        provided.template.body_size_pt = 23
        assert third.config.template.body_size_pt == 11


@pytest.mark.parametrize(
    "argument",
    [
        "-fmarkdown",
        "-tlatex",
        "-oother.docx",
        "--output=other.docx",
        "--from",
        "--",
        "-ddefaults.yaml",
        "--defaults=defaults.yaml",
    ],
)
def test_reserved_pandoc_arguments_cannot_override_pipeline(argument):
    with patch.object(MarkdownToDocxConverter, "_validate_pandoc"):
        converter = MarkdownToDocxConverter()
        with pytest.raises(ConversionError):
            converter._build_pandoc_args(toc=False, toc_depth=3, extra_args=[argument])


def test_real_process_is_reaped_after_timeout(monkeypatch):
    monkeypatch.setattr(
        "markdown2docx.runner.pypandoc.get_pandoc_path", lambda: sys.executable
    )
    with pytest.raises(PandocTimeoutError, match="probe exceeded"):
        run_pandoc(["-c", "import time; time.sleep(30)"], timeout=1, phase="probe")


def test_write_timeout_preserves_existing_document(tmp_path, monkeypatch):
    source = tmp_path / "input.md"
    source.write_text("Important text", encoding="utf-8")
    destination = source.with_suffix(".docx")
    destination.write_bytes(b"previous document")

    def timeout(*args, **kwargs):
        raise PandocTimeoutError("Pandoc write exceeded 1 seconds")

    monkeypatch.setattr("markdown2docx.converter.write_docx", timeout)
    with pytest.raises(PandocTimeoutError):
        MarkdownToDocxConverter().convert(source, destination, offline=True)
    assert destination.read_bytes() == b"previous document"


def test_library_does_not_change_logging_configuration():
    logger = logging.getLogger("markdown2docx.converter")
    before = (list(logger.handlers), logger.level)
    with patch.object(MarkdownToDocxConverter, "_validate_pandoc"):
        MarkdownToDocxConverter()
    assert (logger.handlers, logger.level) == before


@given(st.text(alphabet="abc中文日本한글العربية123 .,!?\u0301\u200d", max_size=100))
def test_grapheme_fragments_preserve_text(text):
    parts = fragments(text, None, text)
    assert "".join(part[1] for part in parts) == text
    import regex

    actual = [cluster for part in parts for cluster in regex.findall(r"\X", part[1])]
    assert actual == regex.findall(r"\X", text)


def test_neutral_scripts_retain_legacy_adjacent_rules():
    assert fragments("...العربية 中文", None, "...العربية 中文")[0][3] == "Arabic"
    assert fragments(". " * 8000, None, "")[0][1] == ". " * 8000


def test_archive_limits_fail_before_xml_loading(tmp_path):
    from zipfile import ZipFile

    path = tmp_path / "huge.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", b"x" * 1024)
    with pytest.raises(ValidationError, match="xml_bytes"):
        validate_archive(path, LimitsConfig(xml_bytes=128))


@pytest.mark.parametrize(
    "location",
    [
        "https://example.com/image.png",
        "file://host/image.png",
        "data:image/png;base64,AAAA",
    ],
)
def test_remote_resources_are_explicitly_rejected(tmp_path, location):
    ast = {"t": "Image", "c": [["", [], []], [], [location, ""]]}
    with pytest.raises(ConversionError):
        validate_resources(ast, tmp_path, offline=True, restricted=False)


def test_restricted_images_cannot_escape_root(tmp_path):
    root = tmp_path / "document"
    root.mkdir()
    (tmp_path / "outside.png").write_bytes(b"image")
    ast = {"t": "Image", "c": [["", [], []], [], ["../outside.png", ""]]}
    with pytest.raises(ConversionError, match="outside"):
        validate_resources(ast, root, offline=False, restricted=True)


@pytest.mark.parametrize(
    "argument",
    [
        "--filter=program",
        "--lua-filter=program",
        "-Lprogram",
        "--defaults=settings.yaml",
    ],
)
def test_restricted_mode_rejects_extension_programs(argument):
    with pytest.raises(ConversionError):
        validate_restricted_arguments([argument])


def test_report_schema_validates_success_and_failure():
    schema = json.loads(
        files("markdown2docx").joinpath("report_schema.json").read_text()
    )
    report = ConversionReport(
        "out.docx", None, [], [], [], preview={"status": "unverified"}
    )
    validate(report.to_dict(), schema)
    validate(
        {
            "status": "error",
            "error": {"code": "INVALID_ARGUMENT", "message": "bad input"},
        },
        schema,
    )
    assert report.schema_version == "1.0"
    assert (
        ConversionReport("out.docx", None, [], [], []).conversion_id
        != report.conversion_id
    )


@pytest.mark.parametrize("mode", [[], ["--ui-lang", "zh"], ["--render"]])
def test_quiet_cli_has_no_success_or_info_output(tmp_path, mode):
    path = tmp_path / "input.md"
    path.write_text("Hello world", encoding="utf-8")
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "markdown2docx.cli",
            str(path),
            "--quiet",
            "--offline",
            *mode,
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout == ""
    assert "INFO" not in result.stderr


def test_serialized_language_properties_are_independently_checked(
    tmp_path, monkeypatch
):
    from markdown2docx.multilingual import MultilingualDocument, W

    source = tmp_path / "input.md"
    source.write_text("[Français]{lang=fr}", encoding="utf-8")
    destination = source.with_suffix(".docx")
    destination.write_bytes(b"old output")
    original = MultilingualDocument._properties

    def corrupt_language(self, run, spec):
        original(self, run, spec)
        language = run.find(W + "rPr/" + W + "lang")
        if language is not None:
            language.set(W + "val", "en")

    monkeypatch.setattr(MultilingualDocument, "_properties", corrupt_language)
    with pytest.raises(ValidationError, match="language mismatch"):
        MarkdownToDocxConverter().convert(source, destination, offline=True)
    assert destination.read_bytes() == b"old output"


def test_restricted_conversion_stages_valid_local_image(tmp_path):
    import base64

    image = tmp_path / "image.png"
    image.write_bytes(
        base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
        )
    )
    source = tmp_path / "input.md"
    source.write_text("![pixel](image.png)\n", encoding="utf-8")
    path = MarkdownToDocxConverter().convert(source, restricted=True, offline=True)
    from zipfile import ZipFile

    with ZipFile(path) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())


def test_image_staging_rejects_oversize_file(tmp_path):
    from markdown2docx.resources import stage_local_resources

    source = tmp_path / "image.png"
    source.write_bytes(b"x" * 100)
    ast = {"t": "Image", "c": [["", [], []], [], [str(source), ""]]}
    destination = tmp_path / "private"
    destination.mkdir()
    with pytest.raises(ConversionError, match="size limit"):
        stage_local_resources(ast, tmp_path, destination, max_bytes=10)


def test_untrusted_pandoc_archive_never_reaches_extraction(tmp_path, monkeypatch):
    import io
    import runpy

    helpers = runpy.run_path(
        str(Path(__file__).resolve().parents[1] / "scripts/install_pandoc.py")
    )
    monkeypatch.setattr(
        "urllib.request.urlopen",
        lambda *args, **kwargs: io.BytesIO(b"tampered archive"),
    )
    with pytest.raises(ValueError, match="SHA-256"):
        helpers["download_checked"](
            "https://github.com/jgm/pandoc/releases/download/3.12/asset.zip",
            tmp_path / "asset.zip",
            "0" * 64,
        )


def test_unwritable_log_file_emits_structured_configuration_error(tmp_path):
    config = tmp_path / "config.toml"
    missing_log = tmp_path / "missing" / "log.txt"
    config.write_text(
        "[logging]\nfile_path = " + json.dumps(str(missing_log)), encoding="utf-8"
    )
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "markdown2docx.cli",
            "--doctor",
            "--json",
            "--config",
            str(config),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 1
    assert json.loads(result.stdout)["error"]["code"] == "INVALID_ARGUMENT"
    assert "Traceback" not in result.stderr


def test_cjk_context_cache_keeps_local_language_overrides(monkeypatch):
    from markdown2docx import multilingual
    from markdown2docx.fonts import FontResolver

    document = multilingual.MultilingualDocument(
        {"meta": {}, "blocks": []},
        FontResolver(offline=True, directories=[]),
        lang=None,
    )
    original = multilingual.cjk_key
    calls = []

    def classify(lang, text):
        calls.append((lang, text))
        return original(lang, text)

    monkeypatch.setattr(multilingual, "cjk_key", classify)
    for _ in range(20):
        assert document._fragments("漢字", None, "漢字かな")[0][0] == "cjk-jp"
        assert document._fragments("漢字", "zh-Hans", "漢字かな")[0][0] == "cjk-sc"
    assert len(calls) == 2


def test_user_marker_named_style_is_preserved(tmp_path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE

    source = tmp_path / "input.md"
    source.write_text('[Keep style]{custom-style="MD2D000000"}', encoding="utf-8")
    reference = tmp_path / "reference.docx"
    from markdown2docx import DocxTemplateManager

    DocxTemplateManager.create_modern_template(reference)
    template = Document(reference)
    template.styles.add_style("MD2D000000", WD_STYLE_TYPE.CHARACTER).font.bold = True
    template.save(reference)
    output = MarkdownToDocxConverter(reference_doc=reference).convert(
        source, offline=True
    )
    result = Document(output)
    assert result.styles["MD2D000000"].font.bold
    assert any(
        run.style.name == "MD2D000000" for p in result.paragraphs for run in p.runs
    )


def test_reference_removes_lowercase_complex_script_theme(tmp_path):
    import zipfile

    from lxml import etree

    from markdown2docx.multilingual import NS

    source = tmp_path / "arabic.md"
    source.write_text("#### عنوان عربي\n", encoding="utf-8")
    output = MarkdownToDocxConverter().convert(source, offline=True)
    with zipfile.ZipFile(output) as archive:
        styles = etree.fromstring(archive.read("word/styles.xml"))
    assert not any(
        name.lower().endswith("theme")
        for fonts in styles.xpath("//w:rFonts", namespaces=NS)
        for name in fonts.attrib
    )


def test_template_cli_json_validates_published_schema(tmp_path):
    output = tmp_path / "template.docx"
    result = subprocess.run(
        [
            sys.executable,
            "-m",
            "markdown2docx.cli",
            "--create-template",
            str(output),
            "--json",
        ],
        capture_output=True,
        text=True,
        check=True,
        timeout=30,
    )
    report = json.loads(result.stdout)
    schema = json.loads(
        files("markdown2docx").joinpath("report_schema.json").read_text()
    )
    validate(report, schema)
    assert report == {
        "status": "success",
        "kind": "template",
        "output_path": str(output),
    }
    assert output.is_file()
