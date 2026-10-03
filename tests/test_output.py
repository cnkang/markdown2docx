"""Regression coverage for output namespace attacks and atomic publication."""

import os
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocumentType

from markdown2docx.config import MarkdownToDocxConfig
from markdown2docx.converter import MarkdownToDocxConverter
from markdown2docx.exceptions import ConversionError, TemplateError, ValidationError
from markdown2docx.output import staged_docx_output
from markdown2docx.templates import DocxTemplateManager


@pytest.fixture(name="render", params=["conversion", "template"])
def render_docx_fixture(request, monkeypatch, tmp_path):
    """Exercise both real entry points with a controlled document writer."""
    source = tmp_path / "input.md"
    source.write_text("# Document")

    def run(destination, before_write=lambda: None, validate=False):
        if request.param == "conversion":
            monkeypatch.setattr(
                MarkdownToDocxConverter, "_validate_pandoc", lambda _: None
            )

            def convert(*_args, outputfile, **_kwargs):
                before_write()
                Document().save(outputfile)

            monkeypatch.setattr(
                "markdown2docx.converter.pypandoc.convert_file", convert
            )
            return MarkdownToDocxConverter().convert(
                source, destination, validate_output=validate
            )
        original_save = DocumentType.save

        def save(document, path):
            before_write()
            original_save(document, path)

        monkeypatch.setattr(DocumentType, "save", save)
        return DocxTemplateManager.create_modern_template(destination)

    return run


@pytest.mark.parametrize("existing", [False, True])
def test_leaf_symlink_never_changes_target(render, tmp_path, existing):
    """Reject live and dangling leaf links without modifying their targets."""
    target = tmp_path / "protected.txt"
    if existing:
        target.write_text("protected")
    output = tmp_path / "report.docx"
    output.symlink_to(target)
    with pytest.raises((ConversionError, TemplateError), match="symlink"):
        render(output)
    assert output.is_symlink()
    assert target.read_text() == "protected" if existing else not target.exists()


@pytest.mark.parametrize("existing", [False, True])
def test_parent_symlink_is_rejected(render, tmp_path, existing):
    """Reject linked ancestors before creating output directories."""
    protected = tmp_path / "protected"
    if existing:
        protected.mkdir()
    alias = tmp_path / "alias"
    alias.symlink_to(protected, target_is_directory=True)
    with pytest.raises((ConversionError, TemplateError)):
        render(alias / "nested" / "report.docx")
    assert not (protected / "nested").exists()


def test_leaf_swap_during_generation_preserves_target(render, tmp_path):
    """A link planted during generation must not redirect the write."""
    protected = tmp_path / "protected.txt"
    protected.write_text("protected")
    output = tmp_path / "report.docx"

    def swap():
        output.symlink_to(protected)

    with pytest.raises((ConversionError, TemplateError), match="symlink"):
        render(output, swap)
    assert protected.read_text() == "protected"
    assert not list(tmp_path.glob(".markdown2docx-*"))


def test_parent_swap_during_generation_cannot_redirect_write(render, tmp_path):
    """A replaced parent cannot change the pinned publication directory."""
    parent = tmp_path / "output"
    parent.mkdir()
    moved = tmp_path / "original-output"
    protected = tmp_path / "protected"
    protected.mkdir()
    victim = protected / "report.docx"
    victim.write_text("protected")
    output = parent / "report.docx"

    def swap():
        parent.rename(moved)
        parent.symlink_to(protected, target_is_directory=True)

    assert render(output, swap) == output
    assert victim.read_text() == "protected"
    assert Document(moved / "report.docx") is not None
    assert not list(moved.glob(".markdown2docx-*"))


def test_nested_directory_creation_and_regular_overwrite(render, tmp_path):
    """Keep regular-file overwrite and validation working."""
    output = tmp_path / "nested" / "directory" / "report.docx"
    output.parent.mkdir(parents=True)
    output.write_text("old")
    assert render(output, validate=True) == output
    assert Document(output) is not None
    assert not list(output.parent.glob(".markdown2docx-*"))


def test_missing_directories_are_created(render, tmp_path):
    """Preserve recursive output directory creation."""
    output = tmp_path / "new" / "nested" / "report.docx"
    assert render(output) == output
    assert Document(output) is not None


def test_symlink_swap_at_atomic_replace_never_follows_target(monkeypatch, tmp_path):
    """A last-instant link is replaced as an entry, not followed."""
    output = tmp_path / "report.docx"
    victim = tmp_path / "protected.txt"
    victim.write_text("protected")
    original_rename = os.rename

    def swap_then_rename(source, destination, **kwargs):
        output.symlink_to(victim)
        original_rename(source, destination, **kwargs)

    with staged_docx_output(output) as stage:
        stage.write_bytes(b"document")
        monkeypatch.setattr(os, "rename", swap_then_rename)
    assert victim.read_text() == "protected"
    assert not output.is_symlink()
    assert output.read_bytes() == b"document"


def test_no_overwrite_is_atomic_at_publication(monkeypatch, tmp_path):
    """Preserve a destination created just before exclusive publication."""
    output = tmp_path / "report.docx"
    original_link = os.link

    def create_then_link(source, destination, **kwargs):
        output.write_bytes(b"concurrent writer")
        original_link(source, destination, **kwargs)

    with pytest.raises(FileExistsError):
        with staged_docx_output(output, overwrite=False) as stage:
            stage.write_bytes(b"document")
            monkeypatch.setattr(os, "link", create_then_link)
    assert output.read_bytes() == b"concurrent writer"
    assert not list(tmp_path.glob(".markdown2docx-*"))


def test_converter_no_overwrite_preserves_existing_file(monkeypatch, tmp_path):
    """Expose no-overwrite errors through the converter API."""
    monkeypatch.setattr(MarkdownToDocxConverter, "_validate_pandoc", lambda _: None)
    source = tmp_path / "input.md"
    source.write_text("# Document")
    output = source.with_suffix(".docx")
    output.write_bytes(b"original")
    config = MarkdownToDocxConfig()
    config.conversion.overwrite_existing = False
    with pytest.raises(ConversionError, match="already exists"):
        MarkdownToDocxConverter(config=config).convert(source)
    assert output.read_bytes() == b"original"


def test_failed_generation_preserves_existing_output(render, tmp_path):
    """Writer failures must leave the previous document unchanged."""
    output = tmp_path / "report.docx"
    output.write_bytes(b"original")

    def fail():
        raise RuntimeError("writer failed")

    with pytest.raises((ConversionError, TemplateError), match="writer failed"):
        render(output, fail)
    assert output.read_bytes() == b"original"
    assert not list(tmp_path.glob(".markdown2docx-*"))


def test_validation_failure_preserves_existing_output(monkeypatch, tmp_path):
    """Do not publish an invalid staged document."""
    monkeypatch.setattr(MarkdownToDocxConverter, "_validate_pandoc", lambda _: None)
    source = tmp_path / "input.md"
    source.write_text("# Document")
    output = source.with_suffix(".docx")
    output.write_bytes(b"original")

    def invalid(*_args, outputfile, **_kwargs):
        Path(outputfile).write_bytes(b"invalid")

    monkeypatch.setattr("markdown2docx.converter.pypandoc.convert_file", invalid)
    with pytest.raises(ValidationError):
        MarkdownToDocxConverter().convert(source, validate_output=True)
    assert output.read_bytes() == b"original"


def test_relative_output_path(monkeypatch, tmp_path):
    """Preserve caller-relative output path semantics."""
    monkeypatch.chdir(tmp_path)
    output = Path("nested/report.docx")
    with staged_docx_output(output, overwrite=False) as stage:
        stage.write_bytes(b"document")
    assert output.read_bytes() == b"document"


def test_unsupported_platform_fails_closed(monkeypatch, tmp_path):
    """Never fall back to race-prone pathname writes."""
    monkeypatch.setattr(os, "supports_dir_fd", set())
    with pytest.raises(OSError, match="not supported"):
        with staged_docx_output(tmp_path / "report.docx"):
            pytest.fail("Writer must not run without secure directory operations")
    assert not (tmp_path / "report.docx").exists()


def test_writer_ignores_untrusted_tmpdir(monkeypatch, tmp_path):
    """An attacker-writable TMPDIR must not hold the writer's private path."""
    temporary = tmp_path / "untrusted-temp"
    temporary.mkdir(mode=0o777)
    monkeypatch.setenv("TMPDIR", str(temporary))
    monkeypatch.setattr("tempfile.tempdir", str(temporary))
    output = tmp_path / "report.docx"
    with staged_docx_output(output) as stage:
        assert temporary not in stage.parents
        stage.write_bytes(b"document")
    assert output.read_bytes() == b"document"
